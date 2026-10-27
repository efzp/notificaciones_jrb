from __future__ import annotations

import math
import posixpath
import re
import warnings
from datetime import date, datetime, time
from io import BytesIO
from numbers import Real
from zipfile import BadZipFile, ZipFile

from defusedxml import ElementTree as SafeET
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
from openpyxl.utils.datetime import from_excel

from .domain import (AppError, FieldSpec, Issue, Original, ParsedRow, ParsedWorkbook,
                     Profile, blank, no_aplica, normalize_header, source_text, sql_length)

MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def parse_date(value, epoch: datetime, minimum: date, maximum: date):
    """Devuelve (date|None, metodo, error). El perfil local es dia/mes/ano."""
    parsed, method, error = None, "", "FECHA_INVALIDA"
    if isinstance(value, datetime):
        parsed, method = value.date(), "NATIVA"
    elif isinstance(value, date):
        parsed, method = value, "NATIVA"
    elif isinstance(value, bool) or isinstance(value, time):
        pass
    else:
        text = str(value).strip()
        if isinstance(value, Real) and math.isfinite(value) and float(value).is_integer():
            text = str(int(value))
        if re.fullmatch(r"\d{8}", text):
            try:
                parsed, method = datetime.strptime(text, "%Y%m%d").date(), "COMPACTA"
            except ValueError:
                pass
        elif isinstance(value, Real) or re.fullmatch(r"\d{4,6}(?:\.\d+)?", text):
            try:
                number = float(value)
                if math.isfinite(number) and 0 <= number <= 2958465:
                    converted = from_excel(number, epoch)
                    if isinstance(converted, datetime):
                        parsed, method = converted.date(), "SERIE_EXCEL"
            except (ValueError, OverflowError, TypeError):
                pass
        else:
            try:
                if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
                    parsed, method = date.fromisoformat(text), "ISO"
                elif re.match(r"^\d{4}-\d{2}-\d{2}[T ]", text):
                    parsed, method = datetime.fromisoformat(text.replace("Z", "+00:00")).date(), "ISO_HORA"
            except ValueError:
                pass
            if parsed is None:
                for fmt in ("%Y/%m/%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y",
                            "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M"):
                    try:
                        parsed, method = datetime.strptime(text, fmt).date(), "LOCAL_DMY" if not fmt.startswith("%Y") else "YMD"
                        break
                    except ValueError:
                        continue
    if parsed is None:
        return None, method, error
    if not minimum <= parsed <= maximum:
        return None, method, "FECHA_FUERA_RANGO"
    return parsed, method, None


def _sheet_extent(archive: ZipFile, title: str, max_rows: int, max_columns: int, profile: Profile):
    """Ignorar dimensiones infladas por formato; inspeccionar celdas fisicas."""
    book = SafeET.fromstring(archive.read("xl/workbook.xml"))
    sheet = next(s for s in book.findall(f"{MAIN}sheets/{MAIN}sheet") if s.attrib["name"] == title)
    relation_id = sheet.attrib[f"{REL}id"]
    relations = SafeET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    relation = next(r for r in relations if r.attrib.get("Id") == relation_id)
    if relation.attrib.get("TargetMode") == "External":
        raise AppError("EXCEL_INVALIDO", "La hoja no puede ser un recurso externo.")
    target = relation.attrib["Target"]
    path = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
    if not path.startswith("xl/") or path not in archive.namelist():
        raise AppError("EXCEL_INVALIDO", "Referencia de hoja incorrecta.")
    empty_shared_strings, known_shared_headers = set(), set()
    known_headers = {normalize_header(name) for name in profile.aliases} | {normalize_header(f.campo) for f in profile.fields}
    if "xl/sharedStrings.xml" in archive.namelist():
        with archive.open("xl/sharedStrings.xml") as stream:
            index = 0
            for _, element in SafeET.iterparse(stream, events=("end",)):
                if element.tag == f"{MAIN}si":
                    text = "".join(t.text or "" for t in element.iter(f"{MAIN}t"))
                    if not text.strip():
                        empty_shared_strings.add(index)
                    if normalize_header(text) in known_headers:
                        known_shared_headers.add(index)
                    index += 1
                    element.clear()
    last_row, last_column, cells = 1, len(profile.fields), 0
    outside_columns, outside_rows, outside_cells = set(), set(), 0
    uncached = set()
    with archive.open(path) as stream:
        for _, element in SafeET.iterparse(stream, events=("end",)):
            if element.tag == f"{MAIN}c":
                cells += 1
                if cells > 500000:
                    raise AppError("EXCEL_GRANDE", "Demasiadas celdas fisicas en la hoja.", 413)
                formula = element.find(f"{MAIN}f")
                value = element.find(f"{MAIN}v")
                inline = element.find(f"{MAIN}is")
                has_value = value is not None and value.text is not None and bool(value.text.strip())
                if has_value and element.attrib.get("t") == "s" and int(value.text) in empty_shared_strings:
                    has_value = False
                has_text = inline is not None and bool("".join(t.text or "" for t in inline.iter(f"{MAIN}t")).strip())
                if has_value or has_text or formula is not None:
                    coordinate = element.attrib.get("r", "")
                    letters, row = coordinate_from_string(coordinate)
                    column = column_index_from_string(letters)
                    if row == 1 and column <= max_columns:
                        last_column = max(last_column, column)
                    if column > last_column:
                        header = "".join(t.text or "" for t in inline.iter(f"{MAIN}t")) if inline is not None else (value.text if value is not None else "")
                        recognized = (element.attrib.get("t") == "s" and has_value and int(value.text) in known_shared_headers) or normalize_header(header) in known_headers
                        if row == 1 and recognized:
                            raise AppError("ENCABEZADO_FUERA_LIMITE", "Un encabezado conocido excede el limite de columnas del parser.", 422)
                        outside_columns.add(column)
                        outside_rows.add(row)
                        outside_cells += 1
                    else:
                        last_row = max(last_row, row)
                        if row > max_rows + 1:
                            raise AppError("EXCEL_GRANDE", "La hoja supera el limite de filas con contenido.", 413)
                    if formula is not None and not has_value and column <= last_column:
                        uncached.add(coordinate)
                element.clear()
            elif element.tag == f"{MAIN}row":
                element.clear()
    outside = {"columnas": len(outside_columns), "celdas_con_datos": outside_cells,
               "filas_afectadas": len(outside_rows),
               "primera_columna": min(outside_columns) if outside_columns else None,
               "ultima_columna": max(outside_columns) if outside_columns else None}
    return last_row, last_column, uncached, outside


def _mapping(headers, rows, profile, epoch, issues):
    aliases = {}
    fields = {f.campo: f for f in profile.fields}
    for alias, canonical in list(profile.aliases.items()) + [(f.campo, f.campo) for f in profile.fields]:
        normalized = normalize_header(alias)
        if canonical not in fields or normalized in aliases and aliases[normalized] != canonical:
            raise AppError("ALIAS_AMBIGUO", "Los alias del perfil no definen una correspondencia unica.", 503)
        aliases[normalized] = canonical
    mapping, duplicate = {}, False
    for position, header in enumerate(headers):
        canonical = aliases.get(normalize_header(header))
        if canonical:
            if canonical in mapping:
                issues.append(Issue("ENCABEZADO_DUPLICADO", "E", "Dos columnas corresponden al mismo campo.", 1, canonical))
                duplicate = True
            else:
                mapping[canonical] = position
    inferred = []
    anchors = len(mapping)
    # Inferencia general de plantilla: al menos 45 anclas con el mismo desplazamiento.
    # No se registra un alias especial para NO APLICA ni para una columna concreta.
    offsets = {position - (fields[name].orden - 1) for name, position in mapping.items()}
    if not duplicate and len(mapping) >= 45 and len(offsets) == 1:
        offset = next(iter(offsets))
        for spec in profile.fields:
            if spec.campo in mapping:
                continue
            position = spec.orden - 1 + offset
            if not 0 <= position < len(headers) or position in mapping.values():
                continue
            evidence = [row[position].value for row in rows if not blank(row[position].value) and not no_aplica(row[position].value)]
            if not evidence:
                continue
            if spec.tipo_dato == "FECHA":
                consistent = all(parse_date(v, epoch, profile.fecha_minima, profile.fecha_maxima)[0] is not None for v in evidence)
            elif spec.tipo_dato == "IDENTIFICADOR":
                consistent = all(not isinstance(v, (date, datetime, bool)) and
                                 (not isinstance(v, Real) or float(v).is_integer()) for v in evidence)
            else:
                consistent = all(not isinstance(v, (date, datetime, bool)) and
                                 (not spec.longitud_maxima or sql_length(str(v)) <= spec.longitud_maxima) for v in evidence)
            if consistent:
                mapping[spec.campo] = position
                inferred.append({"campo": spec.campo, "columna": position + 1,
                                 "encabezado": str(headers[position] or ""), "metodo": "PLANTILLA_Y_VALORES",
                                 "anclas": anchors})
                issues.append(Issue("ENCABEZADO_INFERIDO", "W", "Encabezado recuperado por plantilla alineada y perfil de valores.", 1, spec.campo))
    for spec in profile.fields:
        if spec.campo not in mapping:
            issues.append(Issue("ENCABEZADO_FALTANTE", "E", "Falta una correspondencia inequívoca para el campo.", 1, spec.campo))
    extras = []
    for position in range(len(headers)):
        if position not in mapping.values():
            nonempty = sum(not blank(row[position].value) for row in rows)
            if nonempty or not blank(headers[position]):
                extras.append({"columna": position + 1, "encabezado": str(headers[position] or ""), "celdas_con_datos": nonempty})
                issues.append(Issue("COLUMNA_ADICIONAL", "W", f"Columna adicional {position + 1}; su contenido permanece en el Excel.", 1))
    return mapping, {"mapping": {k: v + 1 for k, v in mapping.items()}, "inferencias": inferred, "columnas_adicionales": extras}


def _parse_row(number, cells, mapping, profile, epoch, uncached, issues):
    values, masks, originals = {}, {}, []

    def original(spec, raw, reason, parser=None, show=True):
        text = source_text(raw)
        if sql_length(text) > 2000:
            issues.append(Issue("ORIGINAL_LARGO", "E", "El valor original excede la longitud permitida.", number, spec.campo))
        else:
            originals.append(Original(spec.campo, text, reason, parser, show))

    for spec in profile.fields:
        cell = cells[mapping[spec.campo]]
        raw = cell.value
        values[spec.campo] = None
        if getattr(cell, "coordinate", "") in uncached:
            issues.append(Issue("FORMULA_SIN_CACHE", "E", "Recalcule y guarde el Excel antes de cargarlo.", number, spec.campo))
            continue
        if blank(raw):
            if spec.obligatorio:
                issues.append(Issue("CAMPO_OBLIGATORIO", "E", "El campo obligatorio esta vacio.", number, spec.campo))
            continue
        if cell.data_type == "e":
            issues.append(Issue("ERROR_EXCEL", "E", "La celda contiene un error de Excel.", number, spec.campo))
            continue
        if no_aplica(raw):
            if spec.obligatorio:
                issues.append(Issue("CAMPO_OBLIGATORIO", "E", "NO APLICA no es valido en una llave obligatoria.", number, spec.campo))
            elif spec.tipo_parte and spec.bit_no_aplica:
                masks[spec.tipo_parte] = masks.get(spec.tipo_parte, 0) | spec.bit_no_aplica
                if raw != "NO APLICA":
                    original(spec, raw, "VARIANTE_NO_APLICA")
            else:
                original(spec, raw, "NORMALIZACION")
            continue
        if spec.tipo_dato == "FECHA":
            parsed, method, error = parse_date(raw, epoch, profile.fecha_minima, profile.fecha_maxima)
            if error:
                original(spec, raw, error, method or None)
                issues.append(Issue(error, "E" if spec.obligatorio else "W", "Fecha invalida o fuera del rango del perfil.", number, spec.campo))
            else:
                values[spec.campo] = parsed
                if method in {"COMPACTA", "SERIE_EXCEL"}:
                    original(spec, raw, "FECHA_RECUPERADA", method, False)
            continue
        text = source_text(raw)
        if spec.tipo_dato == "IDENTIFICADOR":
            if isinstance(raw, (date, datetime, bool)) or isinstance(raw, Real) and (not math.isfinite(raw) or not float(raw).is_integer()):
                issues.append(Issue("IDENTIFICADOR_INVALIDO", "E", "El identificador debe ser texto o un entero exacto.", number, spec.campo))
                continue
            if isinstance(raw, Real):
                if abs(raw) >= 10 ** 15:
                    issues.append(Issue("IDENTIFICADOR_NUMERICO_LARGO", "E", "Guarde el identificador largo como texto para preservar su precision.", number, spec.campo))
                    continue
                text = str(int(raw))
                if re.fullmatch(r"0{2,50}", cell.number_format or ""):
                    text = text.zfill(len(cell.number_format))
            else:
                text = text.strip()
            if "|" in text:
                issues.append(Issue("IDENTIFICADOR_INVALIDO", "E", "El identificador contiene un separador reservado.", number, spec.campo))
                continue
            if text != source_text(raw):
                original(spec, raw, "NORMALIZACION", "IDENTIFICADOR", False)
        if spec.longitud_maxima and sql_length(text) > spec.longitud_maxima:
            issues.append(Issue("LONGITUD_EXCEDIDA", "E", "El valor excede la longitud configurada; no se trunca.", number, spec.campo))
        elif isinstance(raw, (date, datetime, bool)):
            issues.append(Issue("TIPO_INESPERADO", "E", "Tipo de celda incompatible con el campo de texto.", number, spec.campo))
        else:
            values[spec.campo] = text
    return ParsedRow(number, values, masks, originals)


def parse_workbook(content: bytes, profile: Profile, max_rows: int = 2000, max_columns: int = 512) -> ParsedWorkbook:
    profile.validate()
    workbook = None
    try:
        with ZipFile(BytesIO(content)) as archive:
            infos = archive.infolist()
            if (len(infos) > 5000 or len({i.filename for i in infos}) != len(infos)
                    or sum(i.file_size for i in infos) > 100 * 1024 * 1024
                    or any(i.flag_bits & 1 for i in infos)):
                raise AppError("EXCEL_GRANDE", "El contenedor Excel excede los limites permitidos o esta cifrado.", 413)
            with warnings.catch_warnings():
                warnings.simplefilter("ignore", UserWarning)
                workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
            sheet = workbook.worksheets[0]
            if normalize_header(sheet.title) != "hoja1":
                raise AppError("HOJA_NO_SOPORTADA", "El perfil inicial requiere la primera hoja llamada hoja1.", 422)
            last_row, last_column, uncached, outside = _sheet_extent(archive, sheet.title, max_rows, max_columns, profile)
            sheet.reset_dimensions()
            rows = list(sheet.iter_rows(min_row=1, max_row=last_row, max_col=last_column))
            headers = [cell.value for cell in rows[0]]
            data_rows = [(n, row) for n, row in enumerate(rows[1:], 2)
                         if any(not blank(c.value) or getattr(c, "coordinate", "") in uncached for c in row)]
            issues = []
            mapping, structure = _mapping(headers, [r for _, r in data_rows], profile, workbook.epoch, issues)
            structure["columnas_fuera_del_bloque"] = outside
            if outside["celdas_con_datos"]:
                issues.append(Issue("COLUMNAS_FUERA_DEL_BLOQUE", "W", "Hay contenido adicional fuera del bloque de encabezados; se conserva en el Excel de origen."))
            structure.update({"hoja": sheet.title, "fila_encabezados": 1,
                              "hojas_adicionales_omitidas": max(0, len(workbook.worksheets) - 1)})
            if not data_rows:
                issues.append(Issue("ARCHIVO_VACIO", "E", "No se publica un archivo sin filas de datos."))
            parsed = []
            if not any(i.severidad == "E" for i in issues):
                parsed = [_parse_row(n, row, mapping, profile, workbook.epoch, uncached, issues) for n, row in data_rows]
            return ParsedWorkbook(sheet.title, parsed, issues, structure, len(data_rows))
    except AppError:
        raise
    except (BadZipFile, KeyError, ValueError, StopIteration, IndexError, DefusedXmlException, SafeET.ParseError, OSError):
        raise AppError("EXCEL_INVALIDO", "No se pudo interpretar el archivo Excel.", 422) from None
    finally:
        if workbook:
            workbook.close()
