import json
import re
from pathlib import Path
import openpyxl

p = Path(__file__).parent / 'resultados'
rows = json.loads((p/'correos.json').read_text(encoding='utf-8'))
attachments = json.loads((p/'adjuntos.json').read_text(encoding='utf-8'))
original = json.loads((p.parents[1]/'piloto_20/resultados/correos.json').read_text(encoding='utf-8'))
assert [(r['id_piloto'],r['id_pst']) for r in original] == [(r['id_piloto'],r['id_pst']) for r in rows[:20]]
w = openpyxl.load_workbook(p/'piloto_100_correos.xlsx', data_only=False)
assert w['Correos'].max_row == 108 and w['Correos'].max_column == 35
assert w['Adjuntos'].max_row == 247
assert all(len(s.tables) == 1 for s in w)
assert w['Correos']['C9'].is_date
grouped = {}
for values in w['Cuerpos'].iter_rows(min_row=2, values_only=True):
    # Reverse only the documented prefix used to protect formula-like text.
    text = re.sub(r"^'(?=\s*[=+@-])", '', values[4])
    grouped.setdefault(values[0], []).append((values[1], text))
for r in rows:
    joined = ''.join(t[1] for t in sorted(grouped[r['id_piloto']]))
    assert re.sub(r'\s+', '', r['cuerpo_texto']) == re.sub(r'\s+', '', joined), r['id_piloto']
    assert (p/r['cuerpo_archivo']).read_text(encoding='utf-8') == r['cuerpo_texto']
assert len(rows) == len(grouped) == len({r['id_pst'] for r in rows}) == 100
assert sum(r['adjuntos_reportados'] for r in rows) == len(attachments) == 246
assert all(a['id_piloto'] in grouped for a in attachments)
cached = openpyxl.load_workbook(p/'piloto_100_correos.xlsx', data_only=True)
assert cached['Correos']['B4'].value == 100 and cached['Correos']['D4'].value == 246
print('Verificado: 100 cuerpos completos, 20 originales conservados, 246 adjuntos, fechas, tablas y totales.')
print('Segmentos de cuerpo:', w['Cuerpos'].max_row-1)
