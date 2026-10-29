from __future__ import annotations
from pathlib import Path
from collections import Counter, defaultdict
from datetime import date, datetime
import csv, hashlib, importlib.util, json, re
from openpyxl import load_workbook

BASE = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location('sabana_original', r'C:\Users\eduar\Downloads\procesar_reportes_sabana.py')
m = importlib.util.module_from_spec(spec)
spec.loader.exec_module(m)

def json_default(x):
    if isinstance(x, (date, datetime)):
        return x.isoformat()
    if isinstance(x, (set, tuple)):
        return list(x)
    raise TypeError(type(x).__name__)

def source_key(r):
    return (r.get('sala'), re.split(r'[\\/]', str(r.get('archivo_origen', '')))[-1].casefold(), int(r.get('fila_origen', 0)))

wb = load_workbook(BASE/'reporte_sabana_completo.xlsx', read_only=True, data_only=True)
sheet = wb['Consolidado']
rows = sheet.iter_rows(values_only=True)
output_headers = list(next(rows))
reference = [dict(zip(output_headers, r)) for r in rows if any(x is not None for x in r)]
ref_by_source = {source_key(r):r for r in reference}
festivos = {m.parse_fecha(r[0]) for r in list(wb['Festivos'].values)[1:] if m.parse_fecha(r[0])}
reference_errors = list(wb['Errores'].values)
wb.close()

files=[]
header_sets=defaultdict(list)
types=defaultdict(Counter)
specials=defaultdict(Counter)
date_formats=defaultdict(Counter)
invalid_dates=Counter()
invalid_date_locations=[]
extra_column_values=[]
extra_column_seen=set()
extra_column_cell_count=0
max_lengths=Counter()
missing=Counter()
states=defaultdict(Counter)
guides=defaultdict(Counter)
dictamen=Counter()
hashes=defaultdict(list)
all_records=[]
diffs=Counter()
diff_files=Counter()
new_rows=0
equal_rows=0
compared_rows=0
formula_cells=0
for p in sorted((BASE/'inputs').rglob('*')):
    if p.suffix.lower() not in ('.xlsx', '.xlsm'):
        continue
    w=load_workbook(p, read_only=True, data_only=False)
    s=w.worksheets[0]
    it=s.iter_rows()
    header_cells=next(it)
    headers=[m.normalizar_texto(c.value) for c in header_cells]
    header_sets[tuple(headers)].append(f'{p.parent.name}/{p.name}')
    n=0
    duplicate_keys=Counter()
    for cells in it:
        values=[c.value for c in cells]
        if not any(v not in (None,'') for v in values):
            continue
        n+=1
        r=dict(zip(headers,values))
        for c,v in zip(headers,values):
            if v is None or v=='':
                types[c]['empty']+=1
                continue
            types[c][type(v).__name__]+=1
            max_lengths[c]=max(max_lengths[c],len(str(v)))
            if c not in output_headers:
                extra_column_cell_count+=1
                ek=(p.parent.name,p.name,cells[0].row,c,str(v))
                if ek not in extra_column_seen:
                    extra_column_seen.add(ek)
                    extra_column_values.append({'sala':p.parent.name,'archivo':p.name,'fila':cells[0].row,'columna':c,'valor':str(v)})
            if isinstance(v,str) and v.strip().upper() in ('NO APLICA','N/A','NA','NOAPLICA'):
                specials[c][v]+=1
            if c.startswith('fecha_'):
                f=m.parse_fecha(v)
                if not f and not m.es_no_aplica(v):
                    invalid_dates[c]+=1
                    invalid_date_locations.append({'sala':p.parent.name,'archivo':p.name,'fila':cells[0].row,'columna':c,'valor':str(v)})
                if isinstance(v,str):
                    if re.fullmatch(r'\d{4}-\d{2}-\d{2}',v.strip()): fmt='YYYY-MM-DD'
                    elif re.fullmatch(r'\d{4}/\d{2}/\d{2}',v.strip()): fmt='YYYY/MM/DD'
                    elif re.fullmatch(r'\d{2}/\d{2}/\d{4}',v.strip()): fmt='DD/MM/YYYY'
                    else: fmt='otro_texto'
                else: fmt=type(v).__name__
                date_formats[c][fmt]+=1
            if c.startswith('estado_'):
                states[c][str(v)]+=1
            if c.startswith('guia_'):
                guides[c][str(v)]+=1
        for cell in cells:
            if cell.data_type=='f': formula_cells+=1
        for c in ('fecha_audiencia','reparto_paciente','identificacion_paciente'):
            if r.get(c) in (None,''): missing[c]+=1
        dictamen[str(r.get('tipo_de_dictamen'))]+=1
        key=(str(r.get('fecha_audiencia')),str(r.get('reparto_paciente')),str(r.get('identificacion_paciente')))
        duplicate_keys[key]+=1
    files.append({'sala':p.parent.name,'archivo':p.name,'hojas':w.sheetnames,'filas_datos':n,'columnas':len(headers),'filas_extra_por_llave_repetida':sum(v-1 for v in duplicate_keys.values()),'bytes':p.stat().st_size})
    w.close()
    hashes[hashlib.sha256(p.read_bytes()).hexdigest()].append(f'{p.parent.name}/{p.name}')
    records, errors=m.leer_archivo(p, BASE/'inputs', festivos)
    if errors: files[-1]['errores_lectura']=errors
    all_records.extend(records)
    for r in records:
        old=ref_by_source.get(source_key(r))
        if old is None:
            new_rows+=1
            continue
        compared_rows+=1
        changed=[]
        for c in output_headers:
            if c=='archivo_origen':continue
            a=r.get(c)
            b=old.get(c)
            if a in (None,'') and b in (None,''):continue
            if a!=b:
                diffs[c]+=1
                changed.append(c)
        if changed:diff_files[f'{p.parent.name}/{p.name}']+=1
        else:equal_rows+=1

expected_source=[h for h in output_headers if h not in ('sala','archivo_origen','fila_origen','llave_caso') and not h.startswith('dias_') and h!='fecha_primer_envio_partes']
key_counts=Counter(r['llave_caso'] for r in all_records)
ref_file_counts=Counter((r['sala'],re.split(r'[\\/]',str(r['archivo_origen']))[-1]) for r in reference)
source_file_counts={(f['sala'],f['archivo']):f['filas_datos'] for f in files}
inventory=json.loads((BASE/'inventario_sharepoint.json').read_text(encoding='utf-8-sig'))
not_downloaded=[{'sala':i['sala'],'archivo':i['name']} for i in inventory if (i['sala'],i['name']) not in source_file_counts]
result={
 'inventario_archivos':len(inventory),'archivos_revisados':len(files),'pendientes':not_downloaded,
 'por_sala':{s:{'archivos':sum(f['sala']==s for f in files),'filas':sum(f['filas_datos'] for f in files if f['sala']==s)} for s in ('Sala 1','Sala 2','Sala 3')},
 'referencia':{'filas':len(reference),'columnas':len(output_headers),'encabezados':output_headers,'archivos_distintos':len(ref_file_counts),'errores':reference_errors[1:]},
 'esquemas':[{'columnas':len(h),'archivos':len(ff),'encabezados_no_vacios':[x for x in h if x],'cantidad_encabezados_vacios':h.count(''),'ejemplos':ff[:3],'faltantes_vs_referencia':sorted(set(expected_source)-set(h)),'adicionales_vs_referencia':sorted(set(h)-set(expected_source))} for h,ff in header_sets.items()],
 'tipos_por_columna':types,'no_aplica_por_columna':specials,'formatos_fechas':date_formats,'fechas_invalidas_no_vacias':invalid_dates,
 'fechas_invalidas_ubicaciones':invalid_date_locations,'valores_columnas_adicionales':extra_column_values,'celdas_no_vacias_columnas_adicionales':extra_column_cell_count,'longitudes_maximas':max_lengths,
 'llaves_incompletas':missing,'estados':states,'guias':{c:dict(v.most_common(25)) for c,v in guides.items()},'tipos_dictamen':dictamen,
 'formulas_celdas':formula_cells,'total_filas_actuales':len(all_records),'casos_unicos_actuales':len(key_counts),'filas_adicionales_misma_llave':sum(v-1 for v in key_counts.values()),
 'copias_contenido_identico':[v for v in hashes.values() if len(v)>1],
 'comparacion':{'filas_comparadas':compared_rows,'filas_iguales_sin_ruta_local':equal_rows,'filas_nuevas':new_rows,'diferencias_por_columna':diffs,'diferencias_por_archivo':diff_files,
 'archivos_referencia_ausentes_en_actual': [{'sala':s,'archivo':f,'filas':n} for (s,f),n in ref_file_counts.items() if (s,f) not in source_file_counts],
 'conteos_distintos':[{'sala':s,'archivo':f,'actual':n,'referencia':ref_file_counts.get((s,f))} for (s,f),n in source_file_counts.items() if ref_file_counts.get((s,f))!=n]},
 'archivos':files,
}
(BASE/'resultado_estructura.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,default=json_default),encoding='utf-8')
with (BASE/'diccionario_consolidado.csv').open('w',encoding='utf-8-sig',newline='') as fh:
    writer=csv.writer(fh)
    writer.writerow(['orden','columna','procedencia','regla_actual','tipos_observados_entrada'])
    for n,h in enumerate(output_headers,1):
        origin='Excel original' if h in expected_source else ('Calculada' if h.startswith('dias_') or h=='fecha_primer_envio_partes' else 'Metadato / llave')
        if h=='sala':rule='Nombre de subcarpeta Sala N; la funcion debe recibirlo como metadato'
        elif h=='archivo_origen':rule='Ruta del archivo; en Azure se usara ruta SharePoint en lugar de ruta local'
        elif h=='fila_origen':rule='Numero de fila fisica del Excel; comienza en 2'
        elif h=='llave_caso':rule='sala|fecha_audiencia_parseada_ISO|reparto_paciente|identificacion_paciente; no identifica una fila unica'
        elif h=='fecha_primer_envio_partes':rule='Menor fecha valida de envio entre las seis partes; vacio si no existe'
        elif h.startswith('dias_cal_'):rule='Diferencia en dias calendario entre las fechas indicadas; nulo si alguna fecha no es valida'
        elif h.startswith('dias_hab_'):rule='Dias lunes a viernes excluyendo festivos suministrados; excluye inicio e incluye fin; admite signo negativo; nulo si falta fecha valida'
        elif h in ('fecha_audiencia','reparto_paciente','identificacion_paciente'):rule='valor_salida y conversion a texto sin espacios exteriores; no rellena ceros'
        else:rule='valor_salida: fecha reconocida se convierte a ISO; el resto conserva su valor, incluidos NO APLICA y vacios'
        writer.writerow([n,h,origin,rule,json.dumps(types.get(h,{}),ensure_ascii=False)])
summary={k:result[k] for k in ('archivos_revisados','por_sala','total_filas_actuales','casos_unicos_actuales','filas_adicionales_misma_llave','formulas_celdas','fechas_invalidas_no_vacias','llaves_incompletas')}
summary['comparacion']={k:v for k,v in result['comparacion'].items() if k!='archivos_referencia_ausentes_en_actual'}
summary['archivos_referencia_ausentes_en_actual']=len(result['comparacion']['archivos_referencia_ausentes_en_actual'])
print(json.dumps(summary,ensure_ascii=False,default=json_default))
